"""Génère le corpus de démonstration (fichiers PDF, DOCX, TXT et PDF scanné) à partir
des sources Markdown de `sources/`.

Usage :
    python generate_corpus.py [--out ../sample_documents]

Chaque source contient un en-tête YAML (métadonnées + annotations de référence) puis le
corps du document. Les annotations de référence (`gold`) ne sont JAMAIS écrites dans les
fichiers générés : elles servent uniquement à l'évaluation.

Le rendu ajoute en tête de chaque document un « cartouche » (référence, version, date,
direction propriétaire, rédacteur, classification, diffusion), comme dans un référentiel
documentaire bancaire. Le pipeline d'ingestion relit ce cartouche pour extraire les
métadonnées.

Le générateur est déterministe (graine fixe) : deux exécutions produisent les mêmes
contenus textuels, ce qui garantit la reproductibilité des expériences.
"""

from __future__ import annotations

import argparse
import datetime as dt
import random
import re
import shutil
from pathlib import Path

import yaml

HERE = Path(__file__).resolve().parent
SOURCES = HERE / "sources"
REFERENTIALS = HERE / "referentials" / "referentials.yaml"
FONT_REGULAR = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
FONT_SERIF = "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"

DIFFUSION_LABELS = {"ALL": "Tous collaborateurs"}


def load_names() -> dict[str, str]:
    data = yaml.safe_load(REFERENTIALS.read_text(encoding="utf-8"))
    names: dict[str, str] = {}
    for key, items in data.items():
        if key == "ambiguous_aliases":
            continue
        for item in items:
            names[item["id"]] = item["name"]
    return names


def parse_source(path: Path) -> tuple[dict, str]:
    raw = path.read_text(encoding="utf-8")
    match = re.match(r"^---\n(.*?)\n---\n(.*)$", raw, flags=re.S)
    if not match:
        raise ValueError(f"En-tête YAML manquant : {path}")
    return yaml.safe_load(match.group(1)), match.group(2).strip() + "\n"


def cartouche_lines(meta: dict, names: dict[str, str]) -> list[str]:
    date = meta["date"]
    if isinstance(date, (dt.date, dt.datetime)):
        date_txt = date.strftime("%d/%m/%Y")
    else:
        date_txt = str(date)
    diffusion = meta.get("diffusion", ["ALL"])
    if diffusion == ["ALL"]:
        diffusion_txt = "Tous collaborateurs"
    else:
        diffusion_txt = "Restreinte : " + ", ".join(diffusion)
    return [
        "ORIZIA BANQUE — document fictif (corpus de démonstration)",
        f"Référence : {meta['reference']}",
        f"Version : {meta['version']}",
        f"Date : {date_txt}",
        f"Direction propriétaire : {names.get(meta['department'], meta['department'])}",
        f"Rédacteur : {names.get(meta['author'], meta['author'])}",
        f"Classification : {meta['classification']}",
        f"Diffusion : {diffusion_txt}",
    ]


# --------------------------------------------------------------------------- TXT
def render_txt(meta: dict, body: str, names: dict[str, str], out: Path) -> None:
    lines = cartouche_lines(meta, names)
    text = "\n".join(lines) + "\n" + ("=" * 72) + "\n\n" + body
    out.write_text(text, encoding="utf-8")


# --------------------------------------------------------------------------- DOCX
def render_docx(meta: dict, body: str, names: dict[str, str], out: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    table = doc.add_table(rows=0, cols=1)
    table.style = "Table Grid"
    for line in cartouche_lines(meta, names):
        table.add_row().cells[0].text = line
    doc.add_paragraph()

    for block in body.split("\n"):
        stripped = block.rstrip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", stripped):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", stripped), style="List Number")
        else:
            doc.add_paragraph(stripped)

    props = doc.core_properties
    props.title = meta["title"]
    props.author = names.get(meta["author"], meta["author"])
    props.subject = meta["reference"]
    props.keywords = meta["classification"]
    date = meta["date"]
    props.created = dt.datetime(date.year, date.month, date.day, 9, 0, 0)
    props.modified = props.created
    props.last_modified_by = props.author
    doc.save(out)


# --------------------------------------------------------------------------- PDF
def _register_fonts() -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", FONT_REGULAR))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", FONT_BOLD))


def render_pdf(meta: dict, body: str, names: dict[str, str], out: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    _register_fonts()
    base = ParagraphStyle("base", fontName="DejaVu", fontSize=10, leading=14)
    h1 = ParagraphStyle("h1", parent=base, fontName="DejaVu-Bold", fontSize=15, leading=20, spaceAfter=8)
    h2 = ParagraphStyle("h2", parent=base, fontName="DejaVu-Bold", fontSize=12, leading=16, spaceBefore=8, spaceAfter=4)

    story = []
    cart = Table([[Paragraph(escape(line), base)] for line in cartouche_lines(meta, names)], colWidths=[16.5 * cm])
    cart.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke)]))
    story += [cart, Spacer(1, 12)]

    bullets: list = []

    def flush_bullets() -> None:
        if bullets:
            story.append(ListFlowable([ListItem(b) for b in bullets], bulletType="bullet", leftIndent=14))
            bullets.clear()

    for block in body.split("\n"):
        stripped = block.rstrip()
        if not stripped:
            flush_bullets()
            continue
        if stripped.startswith("- "):
            bullets.append(Paragraph(escape(stripped[2:]), base))
            continue
        flush_bullets()
        if stripped.startswith("# "):
            story.append(Paragraph(escape(stripped[2:]), h1))
        elif stripped.startswith("## "):
            story.append(Paragraph(escape(stripped[3:]), h2))
        else:
            story.append(Paragraph(escape(stripped), base))
            story.append(Spacer(1, 4))
    flush_bullets()

    def on_page(canvas, doc_):
        canvas.setFont("DejaVu", 8)
        canvas.drawString(2 * cm, 1.2 * cm, f"{meta['reference']} — v{meta['version']} — {meta['classification']}")
        canvas.drawRightString(19 * cm, 1.2 * cm, f"Page {doc_.page}")

    doc = SimpleDocTemplate(
        str(out), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
        title=meta["title"], author=names.get(meta["author"], meta["author"]), subject=meta["reference"],
        creator="Orizia Banque — corpus fictif", invariant=1,
    )
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


# --------------------------------------------------------------------------- PDF scanné
def render_scanned_pdf(meta: dict, body: str, names: dict[str, str], out: Path, seed: int = 7) -> None:
    """Rend le texte en image (A4, 200 dpi), ajoute du bruit, une légère rotation et un
    flou, puis enregistre un PDF *image seule* (sans couche texte) : seul l'OCR peut en
    extraire le contenu."""
    from PIL import Image, ImageDraw, ImageFilter, ImageFont

    rng = random.Random(seed)
    width, height = 1654, 2339  # A4 à 200 dpi
    img = Image.new("L", (width, height), 250)
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(FONT_SERIF, 30)
    font_bold = ImageFont.truetype(FONT_BOLD, 34)

    y = 120
    for line in cartouche_lines(meta, names):
        draw.text((130, y), line, font=font, fill=20)
        y += 42
    y += 40
    max_chars = 78
    for paragraph in body.split("\n"):
        text = paragraph.strip()
        if not text:
            y += 24
            continue
        is_title = text.isupper() or text.startswith("#")
        text = text.lstrip("# ")
        words, line = text.split(), ""
        for word in words:
            if len(line) + len(word) + 1 > max_chars:
                draw.text((130, y), line, font=font_bold if is_title else font, fill=25)
                y += 42
                line = word
            else:
                line = f"{line} {word}".strip()
        if line:
            draw.text((130, y), line, font=font_bold if is_title else font, fill=25)
            y += 42

    # Dégradations d'un scan réel : bruit poivre et sel, rotation, flou léger.
    pixels = img.load()
    for _ in range(9000):
        x, yy = rng.randrange(width), rng.randrange(height)
        pixels[x, yy] = rng.choice((0, 90, 170))
    img = img.rotate(0.6, resample=Image.BICUBIC, expand=False, fillcolor=250)
    img = img.filter(ImageFilter.GaussianBlur(radius=0.6))
    img.convert("RGB").save(out, "PDF", resolution=200.0)


RENDERERS = {"txt": render_txt, "docx": render_docx, "pdf": render_pdf, "scanned_pdf": render_scanned_pdf}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", default=str(HERE.parent / "sample_documents"))
    args = parser.parse_args()
    out_dir = Path(args.out)
    if out_dir.exists():
        for old in out_dir.iterdir():
            if old.is_file():
                old.unlink()
    out_dir.mkdir(parents=True, exist_ok=True)

    names = load_names()
    sources = sorted(SOURCES.glob("D*.md"))
    produced: dict[str, Path] = {}
    deferred = []
    for src in sources:
        meta, body = parse_source(src)
        if meta.get("copy_of"):
            deferred.append(meta)
            continue
        target = out_dir / meta["file_name"]
        RENDERERS[meta["format"]](meta, body, names, target)
        produced[meta["doc_id"]] = target
    for meta in deferred:
        shutil.copyfile(produced[meta["copy_of"]], out_dir / meta["file_name"])
    print(f"{len(sources)} documents générés dans {out_dir}")


if __name__ == "__main__":
    main()
